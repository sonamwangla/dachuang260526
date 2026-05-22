# -*- coding: utf-8 -*-
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path("project_report.pptx")
FLOW = next(Path.cwd().glob("*.png"), None)
FONT = "Microsoft YaHei"

prs = Presentation()
prs.slide_width = Inches(16)
prs.slide_height = Inches(9)

C = {
    "navy": RGBColor(24, 59, 112),
    "blue": RGBColor(36, 87, 166),
    "green": RGBColor(33, 122, 61),
    "red": RGBColor(180, 35, 24),
    "gold": RGBColor(183, 121, 31),
    "ink": RGBColor(31, 41, 51),
    "muted": RGBColor(82, 96, 109),
    "line": RGBColor(216, 222, 230),
    "bg": RGBColor(244, 246, 248),
    "white": RGBColor(255, 255, 255),
    "pale": RGBColor(237, 242, 247),
}


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def line(shape, color=None, width=1):
    shape.line.color.rgb = color or C["line"]
    shape.line.width = Pt(width)


def text(slide, s, x, y, w, h, size=16, color=None, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = s
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color or C["ink"]
    return box


def bullets(slide, items, x, y, w, h, size=14, color=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color or C["ink"]
        p.space_after = Pt(5)
    return box


def bg(slide, title=None, label=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    fill(shape, C["bg"])
    shape.line.fill.background()
    if title:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.16), prs.slide_height)
        fill(bar, C["blue"])
        bar.line.fill.background()
        text(slide, title, Inches(0.72), Inches(0.35), Inches(10.8), Inches(0.45), 24, C["ink"], True)
        text(slide, label or "Signal Monitoring Console", Inches(11.9), Inches(0.43), Inches(3.2), Inches(0.3), 10, C["muted"], align=PP_ALIGN.RIGHT)
        sep = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(0.96), Inches(14.5), Inches(0.02))
        fill(sep, C["line"])
        sep.line.fill.background()


def card(slide, x, y, w, h, title, body, accent):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    fill(box, C["white"])
    line(box)
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.08), h)
    fill(stripe, accent)
    stripe.line.fill.background()
    text(slide, title, x + Inches(0.24), y + Inches(0.17), w - Inches(0.42), Inches(0.34), 16, C["ink"], True)
    bullets(slide, body, x + Inches(0.24), y + Inches(0.62), w - Inches(0.44), h - Inches(0.75), 12.8, C["muted"])


def metric(slide, label, value, x, y, accent):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.35), Inches(1.15))
    fill(box, C["white"])
    line(box)
    text(slide, value, x + Inches(0.18), y + Inches(0.17), Inches(1.9), Inches(0.38), 24, accent, True)
    text(slide, label, x + Inches(0.18), y + Inches(0.68), Inches(1.9), Inches(0.28), 10.5, C["muted"])


def arrow(slide, x1, y1, x2, y2, color=None):
    conn = slide.shapes.add_connector(1, x1, y1, x2, y2)
    conn.line.color.rgb = color or C["blue"]
    conn.line.width = Pt(2)
    conn.line.end_arrowhead = True


# 1. Cover
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide)
band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.1))
fill(band, C["navy"])
band.line.fill.background()
text(slide, "大学生创新创业训练项目汇报", Inches(0.85), Inches(0.34), Inches(4.8), Inches(0.32), 13, C["white"], True)
text(slide, "西藏高原移动信号\n监测、预测与优化平台", Inches(0.85), Inches(2.0), Inches(9.8), Inches(1.8), 42, C["ink"], True)
text(slide, "面向复杂高原地形的通信覆盖质量分析与运维决策支持系统", Inches(0.9), Inches(4.08), Inches(9.3), Inches(0.45), 20, C["muted"])
bullets(slide, ["拉萨城区及周边采样任务", "RSRP / SINR / 海拔 / GIS 多源融合", "实时监控大屏、3D 仿真、PDF 报告和优化建议"], Inches(0.95), Inches(5.05), Inches(8.0), Inches(1.3), 17)
panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.25), Inches(1.55), Inches(4.6), Inches(5.9))
fill(panel, C["white"])
line(panel)
text(slide, "信号监测闭环", Inches(10.65), Inches(1.95), Inches(3.5), Inches(0.35), 18, C["ink"], True)
for i, (name, col) in enumerate([("采集", C["blue"]), ("建模", C["green"]), ("预测", C["gold"]), ("优化", C["red"])]):
    y = Inches(2.66 + i * 0.82)
    circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.76), y, Inches(0.42), Inches(0.42))
    fill(circ, col)
    circ.line.fill.background()
    text(slide, name, Inches(11.35), y + Inches(0.06), Inches(2.4), Inches(0.25), 14, C["ink"], True)
    if i < 3:
        arrow(slide, Inches(10.97), y + Inches(0.45), Inches(10.97), y + Inches(0.76), col)
text(slide, "项目定位：将通信工程、地理信息、机器学习和前端可视化整合到一套可演示原型中。", Inches(10.65), Inches(6.25), Inches(3.65), Inches(0.75), 12.5, C["muted"])

# 2. Background
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide, "01 项目背景与问题定义", "背景")
card(slide, Inches(0.8), Inches(1.35), Inches(4.55), Inches(2.1), "高原通信挑战", ["高海拔导致传播损耗与气象影响更明显", "山谷、建筑、地形遮挡造成弱覆盖与盲区", "人工排查效率低，难以形成持续监测闭环"], C["red"])
card(slide, Inches(5.75), Inches(1.35), Inches(4.55), Inches(2.1), "平台解决的问题", ["多源采样数据统一存储与管理", "弱覆盖点快速定位和原因分析", "按海拔、天气、运营商维度生成优化建议"], C["blue"])
card(slide, Inches(10.7), Inches(1.35), Inches(4.55), Inches(2.1), "典型应用场景", ["高原城市通信保障", "重点区域巡检与基站规划", "应急通信与网络质量评估"], C["green"])
text(slide, "项目目标", Inches(0.9), Inches(4.1), Inches(2.4), Inches(0.35), 21, C["ink"], True)
bullets(slide, ["构建从采样、建模、预测、可视化到优化处置的完整原型", "用地图热力图、统计图、3D 空间图降低信号问题分析门槛", "为后续真实 GIS/DEM 数据接入和工程化部署打基础"], Inches(0.95), Inches(4.6), Inches(7.9), Inches(1.5), 18)
for i, (label, value, col) in enumerate([("采样点位", "5,000+", C["blue"]), ("平均海拔", "3650m", C["green"]), ("预测准确率", "94.2%", C["gold"]), ("监测周期", "24h", C["red"])]):
    metric(slide, label, value, Inches(9.35 + (i % 2) * 2.62), Inches(4.2 + (i // 2) * 1.45), col)

# 3. Architecture
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide, "02 系统架构与技术栈", "架构")
cols = [
    ("前端展示层", ["Vue3 + Vite + TypeScript", "Vue Router / Axios", "Leaflet 地图与热力图", "ECharts / ECharts-GL 图表"], C["blue"]),
    ("后端服务层", ["Flask API 服务", "JWT 登录鉴权", "SQLAlchemy 数据模型", "上传、历史、统计、预测接口"], C["green"]),
    ("算法分析层", ["MLPRegressor 信号预测", "StandardScaler 特征标准化", "最近基站关联 / Voronoi 思想", "海拔、天气、位置联合建模"], C["gold"]),
    ("数据与运维层", ["SQLite 原型数据库", "模拟器生成实时采样数据", "PDF 报告导出", "后续 Spring Boot + MySQL 部署"], C["red"]),
]
for i, (title, body, col) in enumerate(cols):
    card(slide, Inches(0.8 + i * 3.72), Inches(1.35), Inches(3.35), Inches(4.35), title, body, col)
    if i < 3:
        arrow(slide, Inches(4.05 + i * 3.72), Inches(3.55), Inches(4.35 + i * 3.72), Inches(3.55))
text(slide, "关键接口", Inches(0.9), Inches(6.25), Inches(2.2), Inches(0.3), 18, C["ink"], True)
bullets(slide, ["/api/auth/login：账号登录并返回 token", "/api/upload：上传信号采样点", "/api/history：读取历史采样数据", "/api/stats/comparison：按运营商统计平均 RSRP", "/api/analysis/predict：输入环境参数，返回预测强度、状态与优化建议"], Inches(0.95), Inches(6.65), Inches(13.8), Inches(1.1), 13.8, C["muted"])

# 4. Flow
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide, "03 数据流程与处理闭环", "流程")
if FLOW:
    slide.shapes.add_picture(str(FLOW), Inches(0.95), Inches(1.22), height=Inches(6.92))
card(slide, Inches(6.25), Inches(1.35), Inches(4.1), Inches(1.52), "数据采集层", ["基站信号：RSRP / SINR", "终端信息：设备 / 时间", "GIS：经纬度 / DEM / 地形"], C["blue"])
card(slide, Inches(10.8), Inches(1.35), Inches(4.1), Inches(1.52), "数据处理层", ["清洗、标准化、数据库存储", "距离、海拔差、方位差等特征工程", "多源信号与地理信息融合"], C["green"])
card(slide, Inches(6.25), Inches(3.28), Inches(4.1), Inches(1.52), "算法分析层", ["机器学习模型预测覆盖强度", "识别弱覆盖与高盲区", "输出基站/天线优化策略"], C["gold"])
card(slide, Inches(10.8), Inches(3.28), Inches(4.1), Inches(1.52), "可视化与决策层", ["GIS 地图、热力图、统计图表", "Vue3 交互式监控大屏", "用户操作与运维决策支持"], C["red"])
text(slide, "左侧为项目已有流程图，右侧为答辩讲解时可对应展开的四层闭环。", Inches(6.35), Inches(5.55), Inches(8.3), Inches(0.42), 13.5, C["muted"])

# 5. Algorithm
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide, "04 核心算法与信号预测", "算法")
card(slide, Inches(0.85), Inches(1.25), Inches(4.55), Inches(2.3), "物理模拟模型", ["载波频率：1800 MHz", "基站发射功率：43 dBm", "天线增益：15 dBi", "自由空间损耗 FSPL + 高原海拔修正 + 阴影衰落"], C["blue"])
card(slide, Inches(5.75), Inches(1.25), Inches(4.55), Inches(2.3), "神经网络预测", ["模型：MLPRegressor", "输入：纬度、经度、海拔、天气", "输出：预测 RSRP", "通过 StandardScaler 进行特征标准化"], C["green"])
card(slide, Inches(10.65), Inches(1.25), Inches(4.55), Inches(2.3), "空间覆盖分析", ["根据采样点与基站距离确定最近服务站", "借鉴 Voronoi 空间分区思想", "结合预测阈值生成覆盖状态与优化建议"], C["gold"])
text(slide, "预测结果分级", Inches(0.9), Inches(4.25), Inches(2.2), Inches(0.3), 19, C["ink"], True)
for i, (name, value, col) in enumerate([("优良覆盖", "> -85 dBm", C["green"]), ("良好", "-95 ~ -85 dBm", C["blue"]), ("弱覆盖", "-105 ~ -95 dBm", C["gold"]), ("深度盲区", "< -105 dBm", C["red"])]):
    x = Inches(0.95 + i * 3.55)
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(4.85), Inches(3.05), Inches(1.1))
    fill(box, C["white"])
    line(box)
    text(slide, name, x + Inches(0.18), Inches(5.05), Inches(1.5), Inches(0.25), 15, col, True)
    text(slide, value, x + Inches(0.18), Inches(5.43), Inches(2.3), Inches(0.22), 11.5, C["muted"])
bullets(slide, ["后续扩展：随机森林、XGBoost 用于非线性预测，KMeans 用于弱覆盖区域聚类。", "优化策略：微基站补盲、天线下倾角校准、Massive MIMO 动态波束增强、切换门限优化。"], Inches(0.95), Inches(6.55), Inches(13.8), Inches(0.9), 15)

# 6. Features
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide, "05 功能实现与演示亮点", "功能")
for i, (title, body, col) in enumerate([
    ("账号与权限", ["管理员、操作员、观察员三类角色", "路由级访问控制", "管理员可执行参数调整指令"], C["blue"]),
    ("监控大屏", ["运营商覆盖质量对比", "海拔与 RSRP 衰减散点", "区域信号状态比例", "地图采样点与弱覆盖热力图"], C["green"]),
    ("仿真优化", ["可调海拔与天气环境", "3D Scatter 展示空间分布", "输出预测 RSRP 与处置建议", "支持 PDF 报告导出"], C["gold"]),
]):
    card(slide, Inches(0.95 + i * 4.85), Inches(1.25), Inches(4.3), Inches(2.55), title, body, col)
panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(4.38), Inches(14.15), Inches(3.25))
fill(panel, C["white"])
line(panel)
text(slide, "原型页面结构", Inches(1.25), Inches(4.68), Inches(2.4), Inches(0.3), 17, C["ink"], True)
for label, x, y, w, h, col in [
    ("左侧统计图表", 1.25, 5.28, 3.1, 1.85, C["pale"]),
    ("Leaflet 地图\n采样点 / 热力图 / 基站", 4.65, 5.28, 5.8, 1.85, RGBColor(232, 245, 235)),
    ("右侧模型分析\n环境因子 / 优化建议", 10.75, 5.28, 3.75, 1.85, RGBColor(255, 248, 230)),
]:
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(box, col)
    line(box)
    text(slide, label, Inches(x + 0.18), Inches(y + 0.35), Inches(w - 0.36), Inches(h - 0.55), 15, C["ink"], True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

# 7. Value
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide, "06 创新点与应用价值", "价值")
card(slide, Inches(0.9), Inches(1.3), Inches(4.55), Inches(4.4), "创新点", ["将高原地理环境因素纳入移动信号质量分析", "把物理传播模型与机器学习预测结合", "用 GIS 地图、热力图、3D 空间图呈现覆盖状态", "形成采样、分析、预测、优化建议的闭环"], C["blue"])
card(slide, Inches(5.75), Inches(1.3), Inches(4.55), Inches(4.4), "工程价值", ["缩短弱覆盖问题定位时间", "降低人工巡检和报告整理成本", "为基站补盲、天线调参、网络规划提供依据", "支持运营商网络质量评估与应急保障"], C["green"])
card(slide, Inches(10.6), Inches(1.3), Inches(4.55), Inches(4.4), "可推广场景", ["高原城市与景区通信保障", "铁路、公路、边远区域覆盖巡检", "大型活动临时通信保障", "校园/园区级无线网络质量监测"], C["gold"])
quote = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.35), Inches(6.35), Inches(11.3), Inches(0.9))
fill(quote, C["navy"])
quote.line.fill.background()
text(slide, "让复杂地形下的通信覆盖问题，从“经验判断”变成“数据驱动的可视化决策”。", Inches(2.65), Inches(6.62), Inches(10.7), Inches(0.3), 18, C["white"], True, PP_ALIGN.CENTER)

# 8. Plan
slide = prs.slides.add_slide(prs.slide_layouts[6])
bg(slide, "07 后续计划与阶段目标", "计划")
for i, (phase, title, body, col) in enumerate([
    ("第一阶段", "完善原型演示", ["修复页面中文编码与文案", "补充真实截图与演示数据", "整理接口文档和测试用例"], C["blue"]),
    ("第二阶段", "接入真实数据", ["接入 GIS/DEM 数据", "导入真实采样日志", "建立稳定的数据清洗流程"], C["green"]),
    ("第三阶段", "算法升级", ["引入随机森林、XGBoost", "用 KMeans 聚类弱覆盖区域", "评估预测误差与优化收益"], C["gold"]),
    ("第四阶段", "工程化部署", ["Spring Boot / MySQL 重构规划", "权限、日志、告警能力完善", "性能优化与稳定性测试"], C["red"]),
]):
    x = Inches(0.95 + i * 3.7)
    circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(1.15), Inches(1.35), Inches(0.72), Inches(0.72))
    fill(circ, col)
    circ.line.fill.background()
    text(slide, str(i + 1), x + Inches(1.15), Inches(1.52), Inches(0.72), Inches(0.2), 17, C["white"], True, PP_ALIGN.CENTER)
    if i < 3:
        arrow(slide, x + Inches(1.9), Inches(1.72), x + Inches(3.55), Inches(1.72), col)
    card(slide, x, Inches(2.35), Inches(3.2), Inches(3.55), f"{phase}：{title}", body, col)
text(slide, "预期成果", Inches(0.95), Inches(6.65), Inches(1.8), Inches(0.3), 18, C["ink"], True)
bullets(slide, ["形成可演示的软件原型、项目汇报材料、技术文档和测试数据集。", "具备继续申报竞赛、结题答辩和后续论文/软著整理的基础。"], Inches(2.65), Inches(6.58), Inches(11.8), Inches(0.7), 15)

# 9. End
slide = prs.slides.add_slide(prs.slide_layouts[6])
box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
fill(box, C["navy"])
box.line.fill.background()
text(slide, "谢谢观看", Inches(0.95), Inches(2.25), Inches(5.2), Inches(0.8), 44, C["white"], True)
text(slide, "西藏高原移动信号监测、预测与优化平台", Inches(1.0), Inches(3.45), Inches(7.4), Inches(0.4), 22, RGBColor(218, 228, 240), True)
bullets(slide, ["项目关键词：高原通信、GIS 可视化、机器学习预测、运维优化", "可演示模块：登录权限、监控大屏、3D 仿真、报告导出、优化建议"], Inches(1.05), Inches(4.4), Inches(8.5), Inches(1.1), 17, C["white"])

for idx, s in enumerate(prs.slides, 1):
    if idx not in (1, 9):
        text(s, f"{idx:02d}", Inches(14.65), Inches(8.25), Inches(0.45), Inches(0.22), 10, C["muted"], align=PP_ALIGN.RIGHT)

prs.save(OUT)
print(OUT.resolve())
